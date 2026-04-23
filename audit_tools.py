"""
audit_tools.py - Scans all processed compliance documents in output-documents.zip
for mentions of 8 security/vulnerability tools and generates a comprehensive
Markdown audit report (TOOL_AUDIT_REPORT.md).

Tools audited:
  1. Nessus
  2. Nexpose
  3. Qualys
  4. NMAP
  5. OpenVAS
  6. Malwarebytes
  7. vCenter
  8. PAM360

Usage:
    python3 audit_tools.py
"""

import os
import re
import zipfile
import tempfile
import shutil
from datetime import datetime, timezone
from docx import Document

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

INPUT_ZIP = "output-documents.zip"
REPORT_FILE = "TOOL_AUDIT_REPORT.md"

TOOLS = [
    "Nessus",
    "Nexpose",
    "Qualys",
    "NMAP",
    "OpenVAS",
    "Malwarebytes",
    "vCenter",
    "PAM360",
]

# Number of context characters to capture around each match
CONTEXT_CHARS = 120

# Maximum number of sample excerpts to show per tool per file
MAX_EXCERPTS = 3


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _build_pattern(tool: str) -> re.Pattern:
    """Return a case-sensitive whole-word regex pattern for *tool*."""
    return re.compile(r"\b" + re.escape(tool) + r"\b")


TOOL_PATTERNS = {tool: _build_pattern(tool) for tool in TOOLS}


def extract_text_from_docx(path: str) -> str:
    """Return all plain text from a .docx file (paragraphs + table cells)."""
    doc = Document(path)
    parts: list[str] = []
    for para in doc.paragraphs:
        parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    parts.append(para.text)
    return "\n".join(parts)


def find_occurrences(text: str, pattern: re.Pattern) -> list[dict]:
    """
    Return a list of occurrence dicts with keys:
      - 'pos'     : character offset of the match start
      - 'context' : surrounding text snippet
    """
    results = []
    for m in pattern.finditer(text):
        start = max(0, m.start() - CONTEXT_CHARS)
        end = min(len(text), m.end() + CONTEXT_CHARS)
        snippet = text[start:end].strip().replace("\n", " ")
        results.append({"pos": m.start(), "context": snippet})
    return results


def md_escape(text: str) -> str:
    """Minimal escaping of Markdown special characters in inline text."""
    return text.replace("|", "\\|").replace("*", "\\*").replace("_", "\\_")


# ---------------------------------------------------------------------------
# Core scanning logic
# ---------------------------------------------------------------------------


def scan_documents(extract_dir: str) -> dict:
    """
    Scan every .docx file in *extract_dir* and return a nested dict:

        {
            tool_name: {
                filename: [{"pos": int, "context": str}, ...]
            }
        }
    """
    results: dict[str, dict[str, list]] = {tool: {} for tool in TOOLS}

    doc_files = sorted(
        f for f in os.listdir(extract_dir) if f.lower().endswith(".docx")
    )

    print(f"Scanning {len(doc_files)} documents ...\n")

    for idx, filename in enumerate(doc_files, 1):
        path = os.path.join(extract_dir, filename)
        print(f"  [{idx:3d}/{len(doc_files)}] {filename}")
        try:
            text = extract_text_from_docx(path)
        except Exception as exc:
            print(f"           ✗ Could not read file: {exc}")
            continue

        for tool in TOOLS:
            occurrences = find_occurrences(text, TOOL_PATTERNS[tool])
            if occurrences:
                results[tool][filename] = occurrences

    print()
    return results


# ---------------------------------------------------------------------------
# Report generation
# ---------------------------------------------------------------------------


def generate_report(results: dict, total_files: int) -> str:
    """Build and return the full Markdown report as a string."""

    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")

    lines: list[str] = []

    # ── Header ───────────────────────────────────────────────────────────────
    lines += [
        "# Tool Audit Report",
        "",
        f"> **Generated:** {now}  ",
        f"> **Source archive:** `{INPUT_ZIP}`  ",
        f"> **Total files scanned:** {total_files}",
        "",
        "---",
        "",
    ]

    # ── Summary ──────────────────────────────────────────────────────────────
    lines += [
        "## Summary",
        "",
        "| Tool | Files Mentioning | Total Occurrences |",
        "|------|:---:|:---:|",
    ]
    for tool in TOOLS:
        files_with_tool = results[tool]
        file_count = len(files_with_tool)
        total_occ = sum(len(v) for v in files_with_tool.values())
        lines.append(f"| **{tool}** | {file_count} | {total_occ} |")

    tools_found = sum(1 for tool in TOOLS if results[tool])
    lines += [
        "",
        f"**Tools with at least one mention:** {tools_found} / {len(TOOLS)}",
        "",
        "---",
        "",
    ]

    # ── Detailed Findings ────────────────────────────────────────────────────
    lines += [
        "## Detailed Findings",
        "",
    ]

    for tool in TOOLS:
        files_with_tool = results[tool]
        file_count = len(files_with_tool)
        total_occ = sum(len(v) for v in files_with_tool.values())

        lines += [
            f"### {tool}",
            "",
            f"- **Files mentioning this tool:** {file_count}",
            f"- **Total occurrences across all files:** {total_occ}",
            "",
        ]

        if not files_with_tool:
            lines += ["*No mentions found in any scanned document.*", "", "---", ""]
            continue

        # File-level breakdown
        lines += [
            "| # | File | Occurrences |",
            "|---|------|:---:|",
        ]
        for i, (fname, occs) in enumerate(sorted(files_with_tool.items()), 1):
            lines.append(f"| {i} | `{fname}` | {len(occs)} |")

        # Sample excerpts
        lines += [
            "",
            "#### Sample Excerpts",
            "",
        ]
        for fname, occs in sorted(files_with_tool.items()):
            lines.append(f"**`{fname}`**")
            for j, occ in enumerate(occs[:MAX_EXCERPTS], 1):
                ctx = md_escape(occ["context"])
                lines.append(f"> {j}. …{ctx}…")
            if len(occs) > MAX_EXCERPTS:
                lines.append(
                    f"> *(+ {len(occs) - MAX_EXCERPTS} more occurrence(s) not shown)*"
                )
            lines.append("")

        lines += ["---", ""]

    # ── Cross-Reference ───────────────────────────────────────────────────────
    lines += [
        "## Cross-Reference: Files Mentioning Multiple Tools",
        "",
    ]

    # Build a mapping: filename → set of tools mentioned
    file_to_tools: dict[str, list[str]] = {}
    for tool in TOOLS:
        for fname in results[tool]:
            file_to_tools.setdefault(fname, []).append(tool)

    multi_tool_files = {
        fname: tools for fname, tools in file_to_tools.items() if len(tools) > 1
    }

    if not multi_tool_files:
        lines += [
            "*No files mention more than one of the audited tools.*",
            "",
        ]
    else:
        lines += [
            f"**{len(multi_tool_files)} file(s)** mention more than one audited tool:",
            "",
            "| File | Tools Mentioned | Count |",
            "|------|-----------------|:-----:|",
        ]
        for fname, tools in sorted(
            multi_tool_files.items(), key=lambda x: -len(x[1])
        ):
            tools_str = ", ".join(f"**{t}**" for t in tools)
            lines.append(f"| `{fname}` | {tools_str} | {len(tools)} |")
        lines.append("")

    lines += [
        "---",
        "",
        "*End of report.*",
    ]

    return "\n".join(lines) + "\n"


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------


def main():
    print("=" * 70)
    print("TOOL AUDIT REPORT GENERATOR")
    print("=" * 70)
    print(f"Input archive : {INPUT_ZIP}")
    print(f"Tools audited : {', '.join(TOOLS)}")
    print("=" * 70 + "\n")

    if not os.path.exists(INPUT_ZIP):
        raise FileNotFoundError(
            f"Cannot find '{INPUT_ZIP}'. "
            "Run this script from the repository root."
        )

    with tempfile.TemporaryDirectory() as tmp_dir:
        extract_dir = os.path.join(tmp_dir, "docs")
        os.makedirs(extract_dir)

        # Extract the archive
        print(f"Extracting {INPUT_ZIP} ...\n")
        with zipfile.ZipFile(INPUT_ZIP, "r") as zf:
            for member in zf.namelist():
                # Skip Mac metadata entries
                if "__MACOSX" in member or member.split("/")[-1].startswith("._"):
                    continue
                if member.lower().endswith(".docx"):
                    target = os.path.join(extract_dir, os.path.basename(member))
                    with zf.open(member) as src, open(target, "wb") as dst:
                        shutil.copyfileobj(src, dst)

        doc_files = [
            f for f in os.listdir(extract_dir) if f.lower().endswith(".docx")
        ]
        total_files = len(doc_files)
        print(f"Extracted {total_files} documents.\n")

        # Scan
        results = scan_documents(extract_dir)

    # Generate report
    print("Generating report ...\n")
    report = generate_report(results, total_files)

    with open(REPORT_FILE, "w", encoding="utf-8") as f:
        f.write(report)

    print(f"Report written to: {REPORT_FILE}")
    print()

    # Quick summary to stdout
    print("=" * 70)
    print("AUDIT SUMMARY")
    print("=" * 70)
    for tool in TOOLS:
        file_count = len(results[tool])
        total_occ = sum(len(v) for v in results[tool].values())
        status = f"{file_count} file(s), {total_occ} occurrence(s)"
        print(f"  {tool:<16} : {status}")
    print("=" * 70)


if __name__ == "__main__":
    main()
