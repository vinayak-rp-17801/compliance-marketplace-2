import os
import re
import zipfile
import tempfile
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches

# Input files: split zip batches containing compliance documents
INPUT_ZIPS = [
    "1.zip",
    "2.zip",
    "3.zip",
    "4.zip",
    "5.zip",
]
OUTPUT_ZIP = "output-documents.zip"

# Metadata lines with "Log360 Cloud" branding
METADATA_LINES = [
    "Valid Submission : Yes",
    "Supported Editions : Free, Basic, Standard, Professional, MSSP",
    "Help Document : https://www.manageengine.com/log360-cloud/help/integrations/compliance-extensions.html",
    "Help Video : Not provided",
    "Tags : Log360 Cloud, Compliance, Privacy monitoring, Data privacy, Regulatory compliance",
    "Supported DCs : US, IN, JP, CA, AU, UK and EU",
    "Privacy Policy:",
    "Terms of Service:",
    "Release Notes:",
]

# Pattern to match "Log360" not already followed by " Cloud"
LOG360_STANDALONE = re.compile(r"Log360(?! Cloud)")


def extract_compliance_name(doc_filename):
    """Extract compliance name from a filename like '001_201 CMR 17.00.docx'."""
    name = re.sub(r"^\d+_", "", doc_filename)
    name = name.replace(".docx", "")
    return name


def update_paragraph_runs(paragraph):
    """
    Update runs in a paragraph to replace standalone 'Log360' with 'Log360 Cloud'.

    Handles the edge case where 'Log360 Cloud' is split across two adjacent runs
    (e.g., run ends with 'Log360 ' and the next run starts with 'Cloud'), which
    would otherwise cause a double-replacement producing 'Log360 Cloud Cloud'.
    """
    runs = paragraph.runs
    n = len(runs)
    placeholder = "<<<LOG360_CLOUD_PLACEHOLDER>>>"  # Temporary placeholder for existing "Log360 Cloud"

    for i, run in enumerate(runs):
        text = run.text
        if "Log360" not in text:
            continue

        # Protect existing "Log360 Cloud" already fully present in this run
        protected = text.replace("Log360 Cloud", placeholder)

        # Check if "Log360" at the end of this run merges with the next run
        # to form "Log360 Cloud" across the run boundary.
        next_text = runs[i + 1].text if i + 1 < n else ""
        trailing = re.search(r"Log360\s*$", protected)
        cross_boundary = bool(trailing and re.match(r"\s*Cloud\b", next_text))

        if cross_boundary:
            # The trailing "Log360" belongs to a "Log360 Cloud" split across runs;
            # replace only occurrences that appear before the trailing one.
            last_pos = protected.rfind("Log360")
            prefix = protected[:last_pos]
            suffix = protected[last_pos:]
            prefix = LOG360_STANDALONE.sub("Log360 Cloud", prefix)
            new_text = prefix + suffix
        else:
            new_text = LOG360_STANDALONE.sub("Log360 Cloud", protected)

        run.text = new_text.replace(placeholder, "Log360 Cloud")


def update_metadata_section(doc, compliance_name):
    """
    Replace the metadata section at the end of the document with updated
    'Log360 Cloud' branding. Removes any existing metadata paragraphs
    (starting from 'Valid Submission') and appends fresh ones.
    """
    paragraphs = doc.paragraphs

    # Find the start of the metadata section
    metadata_start = None
    for i, p in enumerate(paragraphs):
        if p.text.startswith("Valid Submission"):
            metadata_start = i
            break

    if metadata_start is not None:
        # Remove existing metadata paragraphs from metadata_start to end
        for i in range(len(paragraphs) - 1, metadata_start - 1, -1):
            p = paragraphs[i]
            p._element.getparent().remove(p._element)

    # Add updated metadata lines
    for line in METADATA_LINES:
        doc.add_paragraph(line)

    # Last line with dynamic compliance name
    doc.add_paragraph(f"Predefined reports for the {compliance_name}.")


def process_document(doc_path, output_path, compliance_name):
    """
    Load a DOCX document, update all 'Log360' references to 'Log360 Cloud'
    in the body text, refresh the metadata section with updated branding,
    and save to output_path.
    """
    doc = Document(doc_path)

    # Determine the boundary of body text vs metadata
    metadata_start = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.startswith("Valid Submission"):
            metadata_start = i
            break

    # Update body text paragraphs only (before the metadata section)
    body_paragraphs = doc.paragraphs[:metadata_start] if metadata_start is not None else doc.paragraphs
    for paragraph in body_paragraphs:
        update_paragraph_runs(paragraph)

    # Update text in tables if any
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    update_paragraph_runs(paragraph)

    # Replace the metadata section with updated "Log360 Cloud" branding
    update_metadata_section(doc, compliance_name)

    doc.save(output_path)


def main():
    print("=" * 70)
    print("STEP 1: Extracting compliance documents from input zip files...")
    print("=" * 70)

    with tempfile.TemporaryDirectory() as tmp_dir:
        docs_extract = os.path.join(tmp_dir, "docs")
        output_docs_dir = os.path.join(tmp_dir, "output_docs")

        os.makedirs(docs_extract, exist_ok=True)
        os.makedirs(output_docs_dir, exist_ok=True)

        # Extract compliance documents from all input zips
        total_extracted = 0
        for zip_name in INPUT_ZIPS:
            if not os.path.exists(zip_name):
                print(f"  WARNING: {zip_name} not found, skipping.")
                continue
            with zipfile.ZipFile(zip_name, "r") as z:
                for item in z.namelist():
                    # Skip Mac metadata and directory entries
                    if "__MACOSX" in item or item.split("/")[-1].startswith("._"):
                        continue
                    if item.endswith(".docx"):
                        target_path = os.path.join(docs_extract, item.split("/")[-1])
                        with z.open(item) as src, open(target_path, "wb") as dst:
                            shutil.copyfileobj(src, dst)
                        total_extracted += 1

        print(f"Extracted {total_extracted} documents.\n")

        print("=" * 70)
        print("STEP 2: Updating 'Log360' references to 'Log360 Cloud'...")
        print("=" * 70 + "\n")

        # Process each document
        doc_files = sorted(
            [f for f in os.listdir(docs_extract) if f.endswith(".docx")]
        )
        total = len(doc_files)
        successful = 0
        failed = 0

        for idx, doc_file in enumerate(doc_files, 1):
            compliance_name = extract_compliance_name(doc_file)
            doc_path = os.path.join(docs_extract, doc_file)
            output_path = os.path.join(output_docs_dir, doc_file)

            print(f"[{idx:3d}/{total}] {doc_file}")
            print(f"         Compliance: {compliance_name}")

            try:
                process_document(doc_path, output_path, compliance_name)
                print(f"         ✓ Updated and saved\n")
                successful += 1
            except Exception as e:
                print(f"         ✗ Error processing document: {e}\n")
                shutil.copy2(doc_path, output_path)
                failed += 1

        # Package output documents into a zip file
        print("=" * 70)
        print("STEP 3: Creating output zip file...")
        print("=" * 70)

        output_files = sorted(os.listdir(output_docs_dir))
        with zipfile.ZipFile(OUTPUT_ZIP, "w", zipfile.ZIP_DEFLATED) as zout:
            for doc_file in output_files:
                file_path = os.path.join(output_docs_dir, doc_file)
                zout.write(file_path, arcname=doc_file)

        print(f"Output written to: {OUTPUT_ZIP}")
        print(f"Total documents in output: {len(output_files)}\n")

    # Summary
    print("=" * 70)
    print("PROCESSING COMPLETE")
    print("=" * 70)
    print(f"Total documents:              {total}")
    print(f"Successfully processed:       {successful}")
    print(f"Errors:                       {failed}")
    print(f"Output zip:                   {OUTPUT_ZIP}")
    print("=" * 70)


if __name__ == "__main__":
    main()
